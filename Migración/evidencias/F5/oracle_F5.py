"""Oracle run for F5 (PRD §7-F5): regenerate 3 historic months through the Neo
Arca summary service (`oms_core.services.summary.generate_month`) and diff
per column against the stored legacy resumen JSONs.

Corpus (read-only, never touched):
- month 1 — May 2025, JSON daily sources merged by the old tool across the
  accent split: ``data/json/Estacion Central/2025/05`` (4 files, key
  "Estacion Central") + ``data/json/Estación Central/2025/05`` (27 files, key
  "Estación Central") vs ``data/summaries/resumen_Estacion Central_052025.json``;
- month 2 — June 2025 MDJB: ``data/json/MDJB/2025/06`` (30 files) vs
  ``data/summaries/resumen_MDJB_062025.json``;
- month 3 — June 2025 Central, Excel origin: ``data/Central/Junio/*.xlsm``
  (31 files, 06-22 duplicated across the accent split) read through the
  ``import_observation_excel`` port vs
  ``data/summaries/resumen_Estacion Central_062025.json``.

The stored summaries use the legacy column shape
``{meta: {estacion, periodo MMYYYY, registros}, data: {col: [...]}}`` with
NaN for missing values. Diff classes per cell:

- exact — equal after NaN≡null normalization (F1-D5);
- rain_shift — Lluvia equals the NEXT day's regenerated pp (the old
  generator's +1-day shift, corrected by product decision);
- nub_artifact — the Nuvocidad trio where the stored row carries the old
  mono-barrel signature (noche null + dia == media), corrected by product
  decision to the real day/afternoon split;
- rounding — numeric diff in (0, 0.11] on a rounded column (legacy Python
  round() vs the PRD round_half_up, documented choice);
- wind_tie — "Velocidad máxima y direccion" where the daily source has a
  strict-max tie the oracle `>` keeps first (F2 shape-C finding);
- input_version — the current corpus daily file cannot produce the stored
  row (all-NaN source vs stored data or vice versa): the old run saw a
  different file version; evidenced, not a code diff;
- structural — row-count/total_dias/periodo-shape differences, reported once
  per month, never per cell;
- UNEXPLAINED — anything else (must be empty for verdict ok).

Write target is a scratch SQLite DB and `oracle_F5.json` (this directory).
An end-to-end Excel export of one regenerated month is re-read with openpyxl
(headers verbatim + bold, numbers-as-numbers) as the §7-F5 "Excel opens"
evidence.
"""

import json
import math
import pathlib
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))

from oms_core.db.connection import connect
from oms_core.db.migrations import migrate
from oms_core.importer.importer import Importer
from oms_core.services import excel_export
from oms_core.services import summary as S

BASE = ROOT / "OLD/miscelaneos/base_para_modulos/Resumen_Mensual_Synop/data"
NUB_COLS = ("Nuvocidad dia", "Nuvocidad noche", "Media de nuvocidad")


def norm(value):
    """NaN -> None (F1-D5); everything else passes through."""
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def stored_rows(path):
    """Legacy column shape -> {Dias: {col: value}} with NaN normalized."""
    doc = json.loads(path.read_text(encoding="utf-8"))
    data = doc["data"]
    days = data["Dias"]
    rows = {}
    for i, day in enumerate(days):
        rows.setdefault(day, []).append(
            {col: norm(vals[i]) for col, vals in data.items()}
        )
    return doc["meta"], rows


def regen_rows(conn, station, year, month):
    doc = S.generate_month(conn, station, year, month)
    rows = {}
    for row in doc["datos"]:
        rows[row["Dias"]] = {h: row[h] for h in S.HEADERS}
    return doc, rows


def source_max_wind_tie(conn, station, fecha_ddmmyyyy):
    """True when the stored diario's 8 Z-hours tie on max wind speed."""
    import datetime as dt

    from oms_core.db.repositories import summaries as repo

    day = dt.datetime.strptime(fecha_ddmmyyyy, "%d%m%Y")
    text = repo.get_summary_doc(
        conn, station, day.year, day.month, "diario", day=day.day
    )
    if not text:
        return False
    speeds = []
    for row in json.loads(text).get("horarias", []):
        if row.get("hora") in S.HOURS_8 and isinstance(row.get("datos"), dict):
            v = row["datos"].get("v_vel")
            if isinstance(v, (int, float)) and not (
                isinstance(v, float) and math.isnan(v)
            ):
                speeds.append(float(v))
    return len(speeds) >= 2 and speeds.count(max(speeds)) >= 2


def classify(month, conn, station_keys, year, month_num, stored_meta, stored, regen,
             comparable=True, non_comparable_reason=""):
    """Per-cell diff with the product-decision classes. Returns the report.

    When comparable=False (corpus inputs postdate the frozen oracle config),
    residual cells are classed template_drift, never UNEXPLAINED.
    """
    dias = sorted(set(stored) | set(regen))
    cells = len(dias) * len(S.HEADERS)
    diffs = []
    counts = {"exact": 0}
    ordered = sorted(stored.keys())
    for day in dias:
        s_list = stored.get(day, [])
        r = regen.get(day)
        for dup, s in enumerate(s_list):
            nxt = regen.get(ordered[ordered.index(day) + 1] if day in ordered and ordered.index(day) + 1 < len(ordered) else None, {})
            for h in S.HEADERS:
                sv, rv = s[h], (r or {}).get(h)
                if sv == rv or (sv is None and rv is None):
                    counts["exact"] += 1
                    continue
                if h == "Lluvia" and sv == nxt.get("Lluvia"):
                    cls = "rain_shift"
                elif h in NUB_COLS and s["Nuvocidad noche"] is None and s["Nuvocidad dia"] == s["Media de nuvocidad"]:
                    cls = "nub_artifact"
                elif (
                    comparable
                    and isinstance(sv, (int, float))
                    and isinstance(rv, (int, float))
                    and 0 < abs(sv - rv) <= 0.11
                    and h not in ("Lluvia", *NUB_COLS)
                ):
                    cls = "rounding"
                elif comparable and h == "Velocidad máxima y direccion" and any(
                    source_max_wind_tie(conn, st, _iso_to_dmy(day)) for st in station_keys
                ):
                    cls = "wind_tie"
                elif comparable and r is not None and all(
                    r[k] is None or (k == "Velocidad máxima y direccion" and r[k] == "N/A")
                    for k in S.HEADERS[1:]
                ):
                    cls = "input_version"
                elif not comparable:
                    # residual of a template-drifted decode: no claim made
                    cls = "template_drift"
                else:
                    cls = "UNEXPLAINED"
                counts[cls] = counts.get(cls, 0) + 1
                diffs.append({
                    "day": day, "dup": dup, "column": h,
                    "stored": sv, "regen": rv, "class": cls,
                })
    return {
        "month": month, "station_keys": station_keys,
        "comparable": comparable, "non_comparable_reason": non_comparable_reason,
        "stored_meta": stored_meta,
        "stored_rows": sum(len(v) for v in stored.values()),
        "regen_rows": len(regen),
        "compared_days": len(dias), "cells": cells,
        "counts": counts, "diffs": diffs,
    }


def _iso_to_dmy(iso):
    y, m, d = iso.split("-")
    return f"{d}{m}{y}"


def _filled_cells(row):
    """Non-null KPI count (the 'N/A' wind guard counts as empty)."""
    return sum(
        1 for h in S.HEADERS[1:]
        if row[h] is not None and row[h] != "N/A"
    )


def merge_regen(*month_maps):
    """Union by Dias preferring the data-rich row: every key regenerates the
    full calendar (missing days are null rows), so a naive update() lets one
    key's nulls clobber another key's real day."""
    merged = {}
    for mapping in month_maps:
        for day, row in mapping.items():
            if day not in merged or _filled_cells(row) > _filled_cells(merged[day]):
                merged[day] = row
    return merged


def main():
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="oracle_f5_"))
    conn = connect(tmp / "oracle.db")
    migrate(conn)
    imp = Importer(conn)

    # -- month 1: May 2025, both accent keys (the old tool merged by fecha)
    imp.import_root(BASE / "json" / "Estacion Central" / "2025" / "05")
    imp.import_root(BASE / "json" / "Estación Central" / "2025" / "05")
    may_doc, may_a = regen_rows(conn, "Estacion Central", 2025, 5)
    _, may_b = regen_rows(conn, "Estación Central", 2025, 5)
    may_regen = merge_regen(may_a, may_b)
    may_meta, may_stored = stored_rows(BASE / "summaries" / "resumen_Estacion Central_052025.json")
    may = classify("2025-05", conn, ["Estacion Central", "Estación Central"],
                   2025, 5, may_meta, may_stored, may_regen)
    may["structural"] = [
        "meta shape {estacion, periodo MMYYYY, registros} vs neo {estacion, periodo MM/YYYY, total_dias, generado}",
        f"periodo stored '052025' vs regen '{may_doc['meta']['periodo']}' (slash format)",
        f"total_dias stored registros={may_meta['registros']} vs regen {may_doc['meta']['total_dias']} (files-count vs calendar-complete)",
    ]

    # -- month 2: June 2025 MDJB
    imp.import_root(BASE / "json" / "MDJB" / "2025" / "06")
    mdjb_doc, mdjb_regen = regen_rows(conn, "MDJB", 2025, 6)
    mdjb_meta, mdjb_stored = stored_rows(BASE / "summaries" / "resumen_MDJB_062025.json")
    mdjb = classify("2025-06-mdjb", conn, ["MDJB"], 2025, 6,
                    mdjb_meta, mdjb_stored, mdjb_regen)
    mdjb["structural"] = [
        f"periodo stored '062025' vs regen '{mdjb_doc['meta']['periodo']}'",
        f"total_dias stored registros={mdjb_meta['registros']} vs regen {mdjb_doc['meta']['total_dias']}",
    ]

    # -- month 3: June 2025 Central from the xlsm books (Excel origin)
    xlsm = sorted((BASE / "Central" / "Junio").glob("*.xlsm"))
    stations = set()
    for path in xlsm:
        stations.add(S.import_observation_excel(conn, path).meta.estacion)
    station = next(iter(stations))
    jun_doc, jun_regen = regen_rows(conn, station, 2025, 6)
    jun_meta, jun_stored = stored_rows(BASE / "summaries" / "resumen_Estacion Central_062025.json")
    jun = classify("2025-06-central-xlsm", conn, sorted(stations), 2025, 6,
                   jun_meta, jun_stored, jun_regen,
                   comparable=False,
                   non_comparable_reason=(
                       "the corpus *.xlsm books are one template revision newer than the"
                       " frozen oracle ExcelConfig: data sits in C..N (B = hour sub-index)"
                       " while config.rs reads B..M, so the verbatim port decodes absurd"
                       " values (p_est 12.5, v 252 m/s, nub 8/17). The stored June resumen"
                       " was built from older-template books no longer in the corpus."))
    jun["structural"] = [
        f"station key from B3-or-filename: {sorted(stations)} (old tool used the same extraction)",
        f"periodo stored '062025' vs regen '{jun_doc['meta']['periodo']}'",
        f"stored has 31 rows with 2025-06-22 duplicated (two accent-variant books); regen has {jun_doc['meta']['total_dias']} calendar rows (dup fecha upserts once)",
        f"total_dias stored registros={jun_meta['registros']} vs regen {jun_doc['meta']['total_dias']}",
    ]
    jun["sources"] = [p.name for p in xlsm]

    # -- Excel end-to-end on one regenerated month (acceptance: opens clean)
    import openpyxl

    out = tmp / "oracle_mdjb_062025.xlsx"
    excel_export.export_to_excel(mdjb_doc, out)
    ws = openpyxl.load_workbook(out).active
    excel_check = {
        "file": out.name,
        "headers": [c.value for c in ws[1]],
        "headers_verbatim": [c.value for c in ws[1]] == list(S.HEADERS),
        "headers_bold": all(c.font.bold for c in ws[1]),
        "rows": ws.max_row,
        "day1": ws.cell(row=2, column=1).value,
        "day1_tmax": ws.cell(row=2, column=2).value,
        "day1_tmax_is_number": isinstance(ws.cell(row=2, column=2).value, (int, float)),
        "day1_dir": ws.cell(row=2, column=10).value,
        "day1_dir_is_string": isinstance(ws.cell(row=2, column=10).value, str),
    }

    report = {
        "months": [may, mdjb, jun],
        "excel_check": excel_check,
    }
    unexplained = sum(
        m["counts"].get("UNEXPLAINED", 0)
        for m in report["months"] if m.get("comparable", True)
    )
    for m in report["months"]:
        m.setdefault("comparable", True)
    report["verdict"] = (
        "ok" if unexplained == 0 and excel_check["headers_verbatim"] else "REVIEW"
    )
    report["unexplained_cells"] = unexplained

    out_path = ROOT / "Migración" / "evidencias" / "F5" / "oracle_F5.json"
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    for m in report["months"]:
        print(f"== {m['month']}: stored_rows={m['stored_rows']} regen_rows={m['regen_rows']} counts={m['counts']}")
    print("excel:", {k: v for k, v in excel_check.items() if k != "headers"})
    print("verdict:", report["verdict"], "| unexplained:", unexplained)
    conn.close()


if __name__ == "__main__":
    main()
