"""summary service — PRD F5: monthly summary generation over SQLite.

Port of the READ-ONLY oracle ``OLD/frontend/src-tauri/src/monthly_summary/``:

- ``json_pipeline.rs`` — the ``DailyObservation``/``DailyKpis`` shapes (the
  21 verbatim serde renames), ``analyze_files`` (adapt flow, adapt of the
  recent synoptic JSON shape, period labeling), ``save_summary``/
  ``load_summary``/``list_summary_history`` (over SQLite here), and the
  Excel surface (ported to ``oms_core.services.excel_export``);
- ``synop_adapter.rs`` — ``adapt_synoptic_json``: the F4 daily document
  shape (``meta + horas``) into the ``DailyObservation`` row model, with
  cloudiness resolution (``Nddff`` digit → ``8NhCLCMCH`` fallback, verbatim)
  and the ``LL_24h`` → ``LL`` precipitation read;
- ``config.rs`` + ``excel_reader.rs`` — the xlsx import mapping (ExcelConfig,
  sheets 3074/4074/1200Z, B3/filename meta, ``get_f64``);
- ``commands.rs`` — the generate/load/history/export/import surface.
  ``calculations.rs`` is NOT reimplemented: per-day KPIs cross the ``oms_calc``
  boundary via ``process_day_data_json`` (F2), already carrying the product
  decisions below.

Product decisions (binding, F5): the OLD generator's bugs are CORRECTED —
rain belongs to its own day (no +1 shift), cloudiness uses the real
day/afternoon split with the documented fallback, NaN is normalized to null
(F1-D5), and rounding follows the PRD ``round_half_up`` (the legacy Python
``round()`` deviation stays a documented, expected oracle diff, quantified
per column in ``Migración/evidencias/F5/``).

Storage (ADR-02): generated documents live in ``summary_logs`` as
``kind='mensual'`` (day=0), never colliding with the imported legacy
``kind='resumen'`` or the daily ``kind='diario'`` documents. One generation
is ONE SQLite transaction (write_guard → transaction, F1 equivalence).

No Qt is imported here (ADR-03); the page lives in ``oms_ui.pages.summary``.
"""

from __future__ import annotations

import calendar
import copy
import json
import math
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import oms_calc
from oms_core.db.repositories import summaries as summaries_repo
from oms_core.services.synoptic import SaveError, _NoCommitConnection, _recalc_derived

#: The 21 VERBATIM column keys of the oracle (json_pipeline serde renames +
#: export order). Do not "fix" the typos: they are the stored contract.
HEADERS: tuple[str, ...] = (
    "Dias",
    "Mayor temperatura máxima",
    "Menor temperatura minima",
    "Media temperatura",
    "MAxima presión nivel medio del mar",
    "Minima presión nivel medio del mar",
    "Media presión nivel medio del mar",
    "Media presion en la estación",
    "Lluvia",
    "Dirección del viento",
    "Velocidad media del viento",
    "Velocidad máxima y direccion",
    "Recorrido del viento",
    "Nuvocidad dia",
    "Nuvocidad noche",
    "Media de nuvocidad",
    "Humedad maxima",
    "Humedad minima",
    "Humedad media",
    "Punto de rocio",
    "Tensión de vapor",
)

#: The 8 synoptic hour labels (synop_adapter hours_config + excel_reader).
HOURS_8: tuple[str, ...] = ("06Z", "09Z", "12Z", "15Z", "18Z", "21Z", "00Z", "03Z")

#: Oracle hours_config: UTC hour <-> AST row index (synop_adapter).
_HOURS_CONFIG: tuple[tuple[str, str], ...] = (
    ("06Z", "2"), ("09Z", "5"), ("12Z", "8"), ("15Z", "11"),
    ("18Z", "14"), ("21Z", "17"), ("00Z", "20"), ("03Z", "23"),
)

#: Oracle temp_config: the hours carrying TMAX/TMIN (<-> cli4074 rows).
_TEMP_CONFIG: tuple[tuple[str, str], ...] = (
    ("06Z", "2"), ("12Z", "8"), ("18Z", "14"), ("00Z", "20"),
)

#: Oracle cloud configs: AST row <-> UTC hour (day 5/8/11, afternoon 14/17/20).
_CLOUD_DAY: tuple[tuple[str, str], ...] = (("5", "09Z"), ("8", "12Z"), ("11", "15Z"))
_CLOUD_AFTERNOON: tuple[tuple[str, str], ...] = (
    (("14", "18Z"), ("17", "21Z"), ("20", "00Z"))
)

#: cli3074 row field -> HourlyData key (synop_adapter cli3074 override read).
_CLI3074_FIELDS: tuple[tuple[str, str], ...] = (
    ("pres_est", "p_est"), ("pres_nmm", "p_nmm"), ("hum_ptor", "rocio"),
    ("hum_tvap", "t_vapor"), ("hum_hr", "hr"),
    ("viento_dir", "v_dir_deg"), ("viento_vel", "v_vel"),
)

#: doc 'horas' raw-input keys that drive derived recalculation (F4 model).
_RECALC_INPUTS = frozenset({
    "ts", "th", "pres_est", "p3", "p24", "correc_alt",
    "meteo_4_irixhvv", "meteo_4_6", "meteo_4_1",
})

#: doc 'horas' derived keys that can be corrected directly: the approved
#: correction IS the new truth over the recomputed value.
_RECALC_OUTPUTS = frozenset({
    "pres_nmm", "punto_rocio", "tension_vapor", "humedad_relativa",
})

#: doc 'horas' datos keys that land on the hourly datos block as strings.
_HORA_DATOS_KEYS = {
    "ts": "ts", "th": "th", "pres_est": "pres_est", "p3": "p3", "p24": "p24",
    "correc_alt": "correc_alt", "viento_dir": "viento_dir",
    "viento_vel": "viento_vel", "visibilidad": "visibilidad",
    # build_horas_map doc keys for the even-hour extras:
    "t_max": "Tmax", "t_min": "Tmin", "t_max_24h": "Tmax_24h",
    "t_min_24h": "Tmin_24h", "ll": "LL", "ll_24h": "LL_24h",
}

#: frontend key ↔ legacy horarias HourlyData key (corrections overlay).
_LEGACY_HORA_DATA_KEYS = {
    "pres_est": "p_est", "pres_nmm": "p_nmm", "punto_rocio": "rocio",
    "tension_vapor": "t_vapor", "humedad_relativa": "hr",
    "viento_dir": "v_dir_deg", "viento_vel": "v_vel",
    "t_max": "tmax", "t_min": "tmin",
}


class SummaryError(Exception):
    """Structured monthly-summary failure ({stage, path, detail})."""

    def __init__(self, detail: str, *, stage: str, path: str) -> None:
        super().__init__(detail)
        self.detail = detail
        self.stage = stage
        self.path = path

    def to_dict(self) -> dict[str, str]:
        return {"stage": self.stage, "path": self.path, "detail": self.detail}


# ---------------------------------------------------------------------------
# row-model shapes (json_pipeline.rs / calculations.rs ports, F2 boundary)
# ---------------------------------------------------------------------------


@dataclass
class HourlyData:
    p_est: float | None = None
    p_nmm: float | None = None
    rocio: float | None = None
    t_vapor: float | None = None
    hr: float | None = None
    v_dir_deg: float | None = None
    v_vel: float | None = None
    tmax: float | None = None
    tmin: float | None = None
    nub: float | None = None


@dataclass
class ObservationMeta:
    estacion: str
    fecha: str  # DDMMYYYY


@dataclass
class HourlyRecord:
    hora: str
    datos: HourlyData


@dataclass
class DaySummary:
    pp: float | None = None


@dataclass
class DailyObservation:
    meta: ObservationMeta
    horarias: list[HourlyRecord] = field(default_factory=list)
    resumen_dia: DaySummary = field(default_factory=DaySummary)


def _num(value: Any) -> float | None:
    """Oracle parse_f64_opt/get_f64 port: float/int/comma-tolerant string.

    NaN becomes None (F1-D5: NaN ≡ NULL everywhere in Neo Arca).
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        cleaned = value.strip().replace(",", ".")
        if not cleaned:
            return None
        try:
            number = float(cleaned)
        except ValueError:
            return None
    else:
        return None
    return None if math.isnan(number) else number


def _first(*values: float | None) -> float | None:
    """First non-null source in oracle priority order."""
    return next((v for v in values if v is not None), None)


def observation_from_rows(
    estacion: str, fecha: str, rows: list[tuple[str, dict]], pp: Any
) -> DailyObservation:
    """Build the typed observation from (hora, datos-dict) rows."""
    allowed = set(_LEGACY_HORA_DATA_KEYS.values())
    pairs: list[tuple[str, dict]] = []
    for hora, datos in rows:
        pairs.append((hora, {k: v for k, v in datos.items() if k in allowed or k == "nub"}))
    records = [HourlyRecord(hora=h, datos=HourlyData(**d)) for h, d in pairs]
    return DailyObservation(
        meta=ObservationMeta(estacion=estacion, fecha=fecha),
        horarias=records,
        resumen_dia=DaySummary(pp=_num(pp)),
    )


def observation_to_dict(obs: DailyObservation) -> dict[str, Any]:
    """Boundary JSON shape of ``oms_calc.process_day_data_json``."""
    return {
        "meta": {"estacion": obs.meta.estacion, "fecha": obs.meta.fecha},
        "horarias": [
            {"hora": row.hora, "datos": vars(row.datos)} for row in obs.horarias
        ],
        "resumen_dia": {"pp": obs.resumen_dia.pp},
    }


def observation_to_kpis(obs: DailyObservation) -> dict[str, Any]:
    """One day through the F2 boundary (the certificado calculation)."""
    payload = json.dumps(observation_to_dict(obs), ensure_ascii=False, allow_nan=False)
    return json.loads(oms_calc.process_day_data_json(payload))


def _empty_kpis(fecha: str) -> dict[str, Any]:
    """Oracle process_day_data on a no-data day: Dias + nulls + 'N/A' guard."""
    return {
        **{header: None for header in HEADERS},
        "Dias": fecha,
        "Velocidad máxima y direccion": "N/A",
    }


# ---------------------------------------------------------------------------
# document adaptation (synop_adapter.rs port over both stored doc shapes)
# ---------------------------------------------------------------------------


def _doc_nub(hora_entry: object) -> float | None:
    """get_nubosidad_from_nddff → get_nubosidad_fallback (verbatim order).

    The oracle read the groups off ``horas[hh]["synop"]``; the schema-2
    documents store the same groups under ``synop[hh]["grupos"]``, so both
    keys are probed and the oracle's order (Nddff first, then 8NhCLCMCH) is
    kept.
    """
    entry = hora_entry if isinstance(hora_entry, dict) else {}
    synop = entry.get("synop")
    if not isinstance(synop, dict):
        synop = entry.get("grupos")
    if not isinstance(synop, dict):
        return None
    if isinstance((nddff := synop.get("Nddff")), str) and nddff and nddff[0].isdigit():
        digit = int(nddff[0])
        if digit <= 9:
            return float(digit)
    g8 = synop.get("8NhCLCMCH")
    if isinstance(g8, str) and len(g8) >= 2 and g8[1].isdigit():
        digit = int(g8[1])
        if digit <= 8:
            return float(digit)
    return None


def _doc_wind_dir(hora_entry: object) -> float | None:
    """Wind direction in degrees from the ``Nddff`` group: digits 2-3 (dd).

    The document does not store ``viento_dir`` (``document_map.DERIVED_DATOS``
    marks it derived): dd is tens of degrees, so '21602' means N=2, dd=16 →
    160 degrees, exactly the value the oracle read from its ``cli3074`` row.
    """
    entry = hora_entry if isinstance(hora_entry, dict) else {}
    synop = entry.get("grupos")
    if not isinstance(synop, dict):
        synop = entry.get("synop")
    if not isinstance(synop, dict):
        return None
    nddff = synop.get("Nddff")
    if not isinstance(nddff, str) or len(nddff) < 3:
        return None
    dd = nddff[1:3]
    if not dd.isdigit():
        return None
    return float(int(dd) * 10)


def _cli_rows(doc: dict[str, Any], key: str) -> dict[str, Any]:
    """The doc's cli block as {AST-row: row-dict}, tolerating int keys.

    Three shapes are read, all present in the real archive:
    - the oracle's flat ``doc["cli3074"]``;
    - the schema-2 ``doc["cli"]["3074"]`` keyed by AST row (the current app's
      storage: only ``viento_vel``/``pres_alti``/``fenomenos``);
    - the migrated legacy block ``doc["cli"]["3074"]["horas"]`` keyed by the 24
      hourly rows, which keeps the full-precision ``temp_seco``/``temp_humedo``
      /``hum_*``/``viento_dir`` values the old app computed with.
    """
    block = doc.get(key)
    if not isinstance(block, dict):
        cli = doc.get("cli")
        block = cli.get(key.removeprefix("cli")) if isinstance(cli, dict) else None
    if not isinstance(block, dict):
        return {}
    horas = block.get("horas")
    if isinstance(horas, dict):
        block = horas
    return {str(row): data for row, data in block.items() if isinstance(data, dict)}


def _cli_num(row: object, field: str) -> float | None:
    """One numeric read off a cli row (NaN normalized to null)."""
    if not isinstance(row, dict):
        return None
    return _num(row.get(field))


def _nub_for_ast(
    row_idx: str, utc_hour: str,
    horas: dict, cli4074: dict[str, Any],
) -> float | None:
    """Oracle get_nubosidad_for_ast_hour: cli4074 row first, then the horas
    UTC slot (Nddff digit, 8NhCLCMCH fallback)."""
    if (nub := _cli_num(cli4074.get(row_idx), "nubosidad")) is not None:
        return nub
    entry = horas.get(utc_hour)
    return _doc_nub(entry) if isinstance(entry, dict) else None


def _fill_derived_from_inputs(
    hourly: dict, datos: dict, *, correc_alt: str = ""
) -> dict:
    """R2: the document no longer stores derived values, so compute them here.

    Feeds the psychrometric pair (ts/th) and the pressure inputs through the
    calculation boundary; fills only what the document did not carry.
    """
    need = [key for key in ("p_nmm", "rocio", "t_vapor", "hr") if hourly.get(key) is None]
    if not need:
        return hourly
    ts = _num(datos.get("ts")) if isinstance(datos, dict) else None
    th = _num(datos.get("th")) if isinstance(datos, dict) else None
    p_est = hourly.get("p_est")
    if p_est is None and isinstance(datos, dict):
        p_est = _num(datos.get("pres_est"))
    if ts is None and th is None and p_est is None:
        return hourly

    def text(value: object) -> str:
        return "" if value is None else str(value)

    try:
        raw = oms_calc.recalculate_derived_fields_json(
            None, text(ts), text(th), text(p_est),
            text(_num(datos.get("p3")) if isinstance(datos, dict) else None),
            "", "", str((datos or {}).get("meteo_4_irixhvv") or ""),
            str((datos or {}).get("meteo_4_6") or ""),
            str((datos or {}).get("meteo_4_1") or ""), "",
            str(correc_alt or ""),
        )
        derived = json.loads(raw) if isinstance(raw, str) else (raw or {})
    except (ValueError, TypeError):
        return hourly
    mapping = {
        "p_nmm": "pres_nmm",
        "rocio": "punto_rocio",
        "t_vapor": "tension_vapor",
        "hr": "humedad_relativa",
    }
    for key, source in mapping.items():
        if hourly.get(key) is not None:
            continue
        value = _num(derived.get(source))
        if value is not None:
            hourly[key] = value
    return hourly


def observation_from_doc(
    doc: dict[str, Any], *, correc_alt: str = ""
) -> DailyObservation | None:
    """Adapt a stored diario document into the ``DailyObservation`` row model.

    Legacy shape ``{meta, horarias, resumen_dia}``: the horarias rows already
    speak the HourlyData key language (p_est/…) and pass through, with NaN
    normalized to null. The F4-revised shape adds ``horas`` — those hour slots
    are the freshest edits, so they win over the stale horarias rows for the 8
    synoptic hours; the cli3074/cli4074 blocks (same doc, F4 save_cli) fill the
    gaps the horas block leaves empty, exactly like the oracle adapter
    (cli4074 rows first for nubosidad and TMAX/TMIN, cli3074 rows as the
    hourly-input fallback). Cloud dia/tarde rows come from the legacy filas
    rows when the doc carries them, else they are derived from the
    cli4074/horas slots per the oracle AST mapping — so a pure F4 day still
    gets its real day/afternoon split instead of the mono-barrel fallback.
    Precipitation: ``resumen_dia.pp`` when stored, else the adapted
    12Z ``LL_24h`` → ``LL`` read (synop_adapter).
    """
    meta = doc.get("meta")
    if not isinstance(meta, dict):
        return None
    estacion, fecha = meta.get("estacion"), meta.get("fecha")
    if not isinstance(estacion, str) or not isinstance(fecha, str):
        return None

    horas = doc.get("horas") if isinstance(doc.get("horas"), dict) else (
        doc.get("synop") if isinstance(doc.get("synop"), dict) else {})
    horarias = [
        row for row in (doc.get("horarias") or [])
        if isinstance(row, dict) and isinstance(row.get("datos"), dict)
    ]
    cli3074 = _cli_rows(doc, "cli3074")
    cli4074 = _cli_rows(doc, "cli4074")
    rows: list[tuple[str, dict]] = []

    def legacy_row(hora: str) -> dict:
        return next((r["datos"] for r in horarias if r.get("hora") == hora), {})

    for hora, row_idx in _HOURS_CONFIG:
        entry = horas.get(hora) if isinstance(horas.get(hora), dict) else None
        legacy = legacy_row(hora)
        calculado = (
            entry.get("calculado") if entry and isinstance(entry.get("calculado"), dict)
            else {}
        )
        datos = entry.get("datos") if entry and isinstance(entry.get("datos"), dict) else {}
        cli = cli3074.get(row_idx)
        hourly = {
            # 'horas' wins where present, then cli3074, then legacy horarias
            "p_est": _first(
                _num(datos.get("pres_est")) if entry else None,
                _cli_num(cli, "pres_est"), _num(legacy.get("p_est")),
            ),
            "p_nmm": _first(
                _num(calculado.get("pres_nmm")) if entry else None,
                _cli_num(cli, "pres_nmm"), _num(legacy.get("p_nmm")),
            ),
            "rocio": _first(
                _num(calculado.get("pr")) if entry else None,
                _cli_num(cli, "hum_ptor"), _num(legacy.get("rocio")),
            ),
            "t_vapor": _first(
                _num(calculado.get("tv")) if entry else None,
                _cli_num(cli, "hum_tvap"), _num(legacy.get("t_vapor")),
            ),
            "hr": _first(
                _num(calculado.get("hr")) if entry else None,
                _cli_num(cli, "hum_hr"), _num(legacy.get("hr")),
            ),
            "v_dir_deg": _first(
                _num(datos.get("viento_dir")) if entry else None,
                _cli_num(cli, "viento_dir"), _doc_wind_dir(entry),
                _num(legacy.get("v_dir_deg")),
            ),
            "v_vel": _first(
                _num(datos.get("viento_vel")) if entry else None,
                _cli_num(cli, "viento_vel"), _num(legacy.get("v_vel")),
            ),
            # Z-row nub only feeds the fallback bucket (filas rows are separate)
            "nub": _first(
                _nub_for_ast(row_idx, hora, horas, cli4074),
                _num(legacy.get("nub")),
            ),
        }
        hourly = _fill_derived_from_inputs(hourly, datos, correc_alt=correc_alt)
        rows.append((hora, hourly))

    # TMAX/TMIN: cli4074 rows first, then the even-hour 'horas' slots, then
    # the legacy rows (the legacy tail only runs on docs without 'horas').
    if horas or cli4074:
        for hora, row_idx in _TEMP_CONFIG:
            slot = horas.get(hora) if isinstance(horas.get(hora), dict) else {}
            slot_datos = slot.get("datos") if isinstance(slot.get("datos"), dict) else {}
            tmax = _first(
                _cli_num(cli4074.get(row_idx), "temp_max"), _num(slot_datos.get("Tmax")),
            )
            tmin = _first(
                _cli_num(cli4074.get(row_idx), "temp_min"), _num(slot_datos.get("Tmin")),
            )
            if tmax is not None:
                rows.append(("TMAX", {"tmax": tmax}))
            if tmin is not None:
                rows.append(("TMIN", {"tmin": tmin}))
    else:
        rows.extend(_aux_rows(horarias, "TMAX", "tmax"))
        rows.extend(_aux_rows(horarias, "TMIN", "tmin"))

    # Cloud buckets: the legacy filas rows win whole-bucket when present (they
    # carry no per-hour labels, so a partial horas edit cannot merge into
    # them); otherwise derive from cli4074/horas per the oracle AST mapping.
    for hora, key in (("filas_nuvosidad_dia", "dia"), ("filas_nuvosidad_tarde", "tarde")):
        legacy = _aux_rows(horarias, hora, "nub")
        if legacy:
            rows.extend(legacy)
            continue
        config = _CLOUD_DAY if key == "dia" else _CLOUD_AFTERNOON
        for row_idx, utc_hour in config:
            if (nub := _nub_for_ast(row_idx, utc_hour, horas, cli4074)) is not None:
                rows.append((hora, {"nub": nub}))

    return observation_from_rows(estacion, fecha, rows, _summary_pp(doc, horas, horarias))


def _aux_rows(horarias: list[dict], hora: str, key: str) -> list[tuple[str, dict]]:
    return [
        (hora, {key: _num(row["datos"].get(key))})
        for row in horarias
        if row.get("hora") == hora
    ]


def _summary_pp(doc: dict, horas: dict, horarias: list[dict]) -> Any:
    """resumen_dia.pp when stored, else the adapted 12Z slot reading.

    Corrected 'll'/'ll_24h' corrections already rewrote ``resumen_dia.pp`` in
    the doc copy (apply_corrections), so they outrank every stored source.
    The stale horarias block of an F4-fused document is passed through to
    keep F4-D2 (meta/horarias/resumen_dia preserved verbatim) intact.
    """
    summary = doc.get("resumen_dia") if isinstance(doc.get("resumen_dia"), dict) else {}
    if "pp" in summary:
        return summary.get("pp")
    entry = horas.get("12Z") if isinstance(horas.get("12Z"), dict) else {}
    datos = entry.get("datos") if entry and isinstance(entry.get("datos"), dict) else {}
    for key in ("LL_24h", "ll_24h", "LL", "ll"):
        if key in datos:
            if (value := _num(datos[key])) is not None:
                return value
            return None
    return None


# ---------------------------------------------------------------------------
# corrections overlay (the oracle *_cor.json workflow, ported to the DB)
# ---------------------------------------------------------------------------


def _parse_correction_value(campo: str, valor: object) -> Any:
    """Corrections carry the oracle 'doc' value language: raw inputs and
    groups as strings, HourlyData fields and precipitation as numbers."""
    if campo in _RECALC_INPUTS or campo in _RECALC_OUTPUTS:
        return "" if str(valor).strip() == "" else str(valor)
    if campo.startswith(("meteo_", "extra_")):
        return "" if str(valor).strip() == "" else str(valor)
    if campo in ("t_max", "t_min", "t_max_24h", "t_min_24h"):
        return "" if str(valor).strip() == "" else str(valor)
    return _num(valor)


def _apply_hour_correction(hora_entry: dict, campo: str, valor: Any) -> bool:
    """Apply an approved correction to one F4 'horas' slot.

    Returns True when the corrected campo changes the compute inputs (so the
    caller must re-run the derived recalculation through oms_calc).
    """
    if campo in _RECALC_OUTPUTS:
        calculado = hora_entry.setdefault("calculado", {})
        doc_key = {
            "pres_nmm": "pres_nmm", "punto_rocio": "pr",
            "tension_vapor": "tv", "humedad_relativa": "hr",
        }[campo]
        _set_or_pop(calculado, doc_key, valor)
        return False
    if campo.startswith(("meteo_", "extra_")):
        synop = hora_entry.setdefault("synop", {})
        _set_or_pop(synop, oms_calc.to_synop_name(campo), valor)
        return campo in _RECALC_INPUTS
    datos = hora_entry.setdefault("datos", {})
    _set_or_pop(datos, _HORA_DATOS_KEYS.get(campo, campo), valor)
    return campo in _RECALC_INPUTS


def _set_or_pop(target: dict, key: str, valor: Any) -> None:
    if valor == "" or valor is None:
        target.pop(key, None)
    else:
        target[key] = valor


def _apply_legacy_correction(
    doc: dict[str, Any], hora: str, campo: str, valor: Any
) -> None:
    """Corrections on a legacy-shaped slot: HourlyData fields only. Raw inputs
    ts/th/p3/p24 are NOT stored in legacy documents (F2 shape-B finding), so
    recomputing from them would invent data — the correction targets the
    stored value instead. When the doc never captured that hour (an F4-fresh
    day without 'horarias'), the row is materialized so the approved fix
    lands instead of being silently dropped."""
    mapped = _LEGACY_HORA_DATA_KEYS.get(campo)
    if mapped is None:
        return
    horarias = doc.get("horarias")
    if not isinstance(horarias, list):
        horarias = doc["horarias"] = []
    for row in horarias:
        if row.get("hora") == hora and isinstance(row.get("datos"), dict):
            _set_or_pop(row["datos"], mapped, _num(valor))
            return  # the slot matches the first row for that hour only
    if (number := _num(valor)) is not None:
        horarias.append({"hora": hora, "datos": {mapped: number}})


def apply_corrections(
    doc: dict[str, Any],
    conn: sqlite3.Connection,
    station: str,
    fecha_iso: str,
    station_ch: str = "",
) -> tuple[dict[str, Any], int]:
    """Overlay the approved audit corrections onto a copy of the daily doc.

    Returns (doc_copy, applied_count). Corrections land on the CURRENT frame
    of the document: F4 'horas' slots re-derive via oms_calc when the corrected
    campo is a raw input; legacy docs accept stored-value corrections only.
    """
    corrections = synoptic_svc_corrections(conn, station, fecha_iso)
    if not corrections:
        return doc, 0
    doc = copy.deepcopy(doc)
    applied = 0
    recompute: dict[str, dict] = {}
    horas = doc.get("horas") if isinstance(doc.get("horas"), dict) else (
        doc.get("synop") if isinstance(doc.get("synop"), dict) else {})
    for hora, campo, valor in corrections:
        value = _parse_correction_value(campo, valor)
        entry = horas.get(hora)
        if isinstance(entry, dict):
            if _apply_hour_correction(entry, campo, value):
                recompute[hora] = entry
        else:
            _apply_legacy_correction(doc, hora, campo, value)
        if campo in ("ll", "ll_24h"):
            # precipitation correction: authority over the day's pp
            summary = doc.setdefault("resumen_dia", {})
            _set_or_pop(summary, "pp", _num(value))
        applied += 1

    for hora, entry in recompute.items():
        _recompute_derived(doc, horas, hora, entry, station_ch)
    return doc, applied


def synoptic_svc_corrections(
    conn: sqlite3.Connection, station: str, fecha_iso: str
) -> list[tuple[str, str, str]]:
    """The approved corrections for the date (audit repo query)."""
    rows = conn.execute(
        "SELECT hora, campo, valor_corregido FROM corrections"
        " WHERE station_id = ? AND fecha = ? AND estado = 'aprobado'",
        (station, fecha_iso),
    ).fetchall()
    return [(row["hora"], row["campo"], row["valor_corregido"]) for row in rows]


def _recompute_derived(
    doc: dict, horas: dict, hora: str, entry: dict, station_ch: str
) -> None:
    """F4 model: raw inputs changed → the derived block is recomputed with
    the certified calculation (``recalculate_derived_fields_json`` through the
    F2 boundary; the station 'ch' is resolved once per generation)."""
    datos = entry.get("datos") if isinstance(entry.get("datos"), dict) else {}
    frontend = {
        "ts": datos.get("ts"), "th": datos.get("th"),
        "pres_est": datos.get("pres_est"), "p3": datos.get("p3"),
        "p24": datos.get("p24"), "correc_alt": datos.get("correc_alt"),
        "viento_dir": datos.get("viento_dir"),
        "viento_vel": datos.get("viento_vel"),
        "visibilidad": datos.get("visibilidad"),
    }
    synop = entry.get("synop") if isinstance(entry.get("synop"), dict) else {}
    for group, campo in (("IrIXHVV", "meteo_4_irixhvv"), ("7wwW1W2", "meteo_4_6"),
                         ("Nddff", "meteo_4_1")):
        if isinstance(synop.get(group), str):
            frontend[campo] = synop[group]
    prev_nddff = ""
    prev = horas.get(oms_calc.get_prev_hour(hora))
    if isinstance(prev, dict) and isinstance(prev.get("synop"), dict):
        value = prev["synop"].get("Nddff")
        if isinstance(value, str):
            prev_nddff = value
    derived = _recalc_derived(frontend, None, prev_nddff, station_ch)
    calculado = entry.setdefault("calculado", {})
    calculado["pres_nmm"] = derived.get("pres_nmm") or None
    calculado["pr"] = derived.get("punto_rocio") or None
    calculado["tv"] = derived.get("tension_vapor") or None
    calculado["hr"] = derived.get("humedad_relativa") or None


# ---------------------------------------------------------------------------
# generation
# ---------------------------------------------------------------------------


def _station_ch(conn: sqlite3.Connection, station: str) -> str:
    row = conn.execute("SELECT ch FROM stations WHERE id = ?", (station,)).fetchone()
    return "" if row is None else str(row["ch"])


def generate_month(
    conn: sqlite3.Connection, station: str, year: int, month: int,
    *, overlay: bool = True,
) -> dict[str, Any]:
    """Generate the monthly summary document for (station, month, year).

    Reads the verbatim 'diario' documents from summary_logs, applies the
    approved audit corrections per day (``overlay=False`` skips them — the
    UI's 'Incluir correcciones (overlay)' switch), recomputes derived inputs
    where they changed, and runs every day through ``process_day_data_json``
    (F2 boundary). The document is stored as ``kind='mensual'`` (day=0), never
    colliding with the imported legacy ``kind='resumen'`` docs; one generation
    is one transaction.
    """
    station = str(station)
    # Monthly-generation validation is station-KEY based (the corpus keys
    # carry names like 'Estación Central'): no empties, no path escapes.
    if not station or "/" in station or "\\" in station:
        raise SummaryError(
            "Código de estación inválido", stage="validate", path=station
        )
    year, month = int(year), int(month)
    if not 1 <= month <= 12:
        raise SummaryError(f"mes inválido: {month}", stage="validate", path=station)
    last_day = calendar.monthrange(year, month)[1]
    station_ch = _station_ch(conn, station)

    per_day: list[dict[str, Any]] = []
    for day in range(1, last_day + 1):
        fecha_iso = f"{year:04d}-{month:02d}-{day:02d}"
        text = summaries_repo.get_summary_doc(conn, station, year, month, "diario", day=day)
        if text is None:
            per_day.append(_empty_kpis(fecha_iso))
            continue
        try:
            doc = json.loads(text)
        except ValueError as exc:
            raise SummaryError(
                f"el documento diario de {station} {fecha_iso} no es JSON válido: {exc}",
                stage="load", path=f"{station}/{year:04d}/{month:02d}/{day:02d}",
            ) from exc
        if not isinstance(doc, dict):
            per_day.append(_empty_kpis(fecha_iso))
            continue
        doc, _applied = (
            apply_corrections(doc, conn, station, fecha_iso, station_ch)
            if overlay else (doc, [])
        )
        obs = observation_from_doc(doc, correc_alt=station_ch)
        per_day.append(observation_to_kpis(obs) if obs else _empty_kpis(fecha_iso))

    doc_data = {
        "meta": {
            "estacion": station,
            "periodo": f"{month:02d}/{year:04d}",
            "total_dias": len(per_day),
            "generado": datetime.now().astimezone().isoformat(),
        },
        "datos": [{header: k.get(header) for header in HEADERS} for k in per_day],
    }

    text = json.dumps(doc_data, ensure_ascii=False)
    sha = _sha256_text(text)
    conn.execute("BEGIN IMMEDIATE")
    try:
        tx = _NoCommitConnection(conn)
        summaries_repo.upsert_summary_doc(
            tx, station_id=station, year=year, month=month,
            kind="mensual", data=text, sha256=sha, day=0,
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return doc_data


# ---------------------------------------------------------------------------
# corrections indicator (UI overlay source)
# ---------------------------------------------------------------------------


def corrections_by_day(
    conn: sqlite3.Connection, station: str, year: int, month: int
) -> dict[int, int]:
    """Approved-correction count per calendar day (the UI overlay source)."""
    prefix = f"{year:04d}-{month:02d}-"
    rows = conn.execute(
        "SELECT fecha, COUNT(*) FROM corrections WHERE station_id = ?"
        " AND estado = 'aprobado' GROUP BY fecha",
        (station,),
    ).fetchall()
    return {
        int(row["fecha"][len(prefix):]): int(row[1])
        for row in rows
        if row["fecha"].startswith(prefix)
    }


# ---------------------------------------------------------------------------
# history (list_summary_history / load / delete ports)
# ---------------------------------------------------------------------------


def list_history(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    """Generated summaries index, newest period first (oracle ms_list_history)."""
    rows = conn.execute(
        "SELECT station_id, year, month, data FROM summary_logs"
        " WHERE kind = 'mensual' ORDER BY year DESC, month DESC, station_id"
    ).fetchall()
    entries: list[dict[str, Any]] = []
    for row in rows:
        try:
            doc = json.loads(row["data"])
        except ValueError:
            continue
        meta = doc.get("meta") if isinstance(doc.get("meta"), dict) else {}
        entries.append({
            "station": row["station_id"],
            "period": f"{row['month']:02d}/{row['year']:04d}",
            "year": int(row["year"]),
            "month": int(row["month"]),
            "total_dias": meta.get("total_dias", len(doc.get("datos") or [])),
            "generado": meta.get("generado", ""),
        })
    return entries


def load_doc(
    conn: sqlite3.Connection, station: str, year: int, month: int
) -> dict[str, Any] | None:
    """load_summary port over the DB document (None when absent)."""
    text = summaries_repo.get_summary_doc(conn, station, int(year), int(month), "mensual")
    if text is None:
        return None
    doc = json.loads(text)
    return doc if isinstance(doc, dict) else None


def delete_doc(conn: sqlite3.Connection, station: str, year: int, month: int) -> bool:
    """ms_delete_summary port (DB row instead of the OS trash file)."""
    return summaries_repo.delete_doc(conn, str(station), int(year), int(month), "mensual")


def get_diario_doc(
    conn: sqlite3.Connection, station: str, year: int, month: int, day: int
) -> str | None:
    """Expose the stored daily document (import-Excel verification + tests)."""
    return summaries_repo.get_summary_doc(conn, station, year, month, "diario", day=day)


def _sha256_text(text: str) -> str:
    import hashlib

    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# xlsx import (excel_reader.rs port)
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ExcelConfig:
    """excel_reader defaults (config.rs get_default_config), 0-indexed."""

    sheet_daily: str = "3074"
    sheet_cloud_temp: str = "4074"
    sheet_rain: str = "1200Z"
    rain_cell: tuple[int, int] = (36, 6)  # G37
    rows_8_observations: tuple[int, ...] = (10, 13, 16, 19, 22, 25, 28, 31)
    rows_cloud_day: tuple[int, ...] = (13, 16, 19)
    rows_cloud_afternoon: tuple[int, ...] = (22, 25, 28)
    rows_4_temperatures: tuple[int, ...] = (10, 16, 22, 28)

    col_presion_estacion: int = 1
    col_presion_nmm: int = 2
    col_punto_rocio: int = 8
    col_tension_vapor: int = 9
    col_humedad_relativa: int = 10
    col_viento_direccion: int = 11
    col_viento_velocidad: int = 12
    col_nuvosidad: int = 1
    col_temp_maxima: int = 21
    col_temp_minima: int = 22


def get_default_config() -> ExcelConfig:
    return ExcelConfig()


_DATE_RE = re.compile(r"(\d{8})")
_STATION_RE = re.compile(r"^(.*?)(?:\s*\d{8})")


def _extract_date_from_filename(path: Path) -> str:
    match = _DATE_RE.search(path.stem)
    return match.group(1) if match else "01011970"


def _extract_station_from_filename(path: Path) -> str:
    match = _STATION_RE.search(path.stem)
    station = match.group(1).strip() if match else ""
    return station or "Estacion"


def _sheet(ws: object) -> Any:
    """calamine Range equivalent (a worksheet or empty)."""
    return ws


def import_observation_excel(conn: sqlite3.Connection, path: str | Path) -> DailyObservation:
    """build_observation_from_excel port: xlsx → DailyObservation + diario doc.

    Mirrors excel_reader.rs: the 8 observation rows of sheet 3074, TMAX/TMIN
    and the cloudiness rows of 4074 (empty when the sheet is missing), the
    1200Z rain cell, the date from the filename (8 digits) and the station
    from B3 falling back to the filename. The adapted daily document is
    stored verbatim-shape as ``kind='diario'`` so the generation sees it.
    """
    import openpyxl

    config = get_default_config()
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise SummaryError(
            f"Failed to open excel: {exc}", stage="import", path=str(path)
        ) from exc

    def cell(sheet: str, row: int, col: int) -> Any:
        if sheet in wb.sheetnames:
            return wb[sheet].cell(row=row + 1, column=col + 1).value
        return None  # unwrap_or_else(Range::empty)

    sheet_daily = config.sheet_daily
    fecha = _extract_date_from_filename(path)
    estacion = _extract_station_from_filename(path)
    b3 = cell(sheet_daily, 2, 1)
    if isinstance(b3, str) and b3.strip():
        estacion = b3.strip()

    horarias: list[dict[str, Any]] = []
    for hora, row in zip(HOURS_8, config.rows_8_observations):
        horarias.append({
            "hora": hora,
            "datos": {
                "p_est": _num(cell(sheet_daily, row, config.col_presion_estacion)),
                "p_nmm": _num(cell(sheet_daily, row, config.col_presion_nmm)),
                "rocio": _num(cell(sheet_daily, row, config.col_punto_rocio)),
                "t_vapor": _num(cell(sheet_daily, row, config.col_tension_vapor)),
                "hr": _num(cell(sheet_daily, row, config.col_humedad_relativa)),
                "v_dir_deg": _num(cell(sheet_daily, row, config.col_viento_direccion)),
                "v_vel": _num(cell(sheet_daily, row, config.col_viento_velocidad)),
            },
        })
    sheet_cloud = config.sheet_cloud_temp
    for row in config.rows_4_temperatures:
        tmax = _num(cell(sheet_cloud, row, config.col_temp_maxima))
        if tmax is not None:
            horarias.append({"hora": "TMAX", "datos": {"tmax": tmax}})
        tmin = _num(cell(sheet_cloud, row, config.col_temp_minima))
        if tmin is not None:
            horarias.append({"hora": "TMIN", "datos": {"tmin": tmin}})
    for hora, rows in (("filas_nuvosidad_dia", config.rows_cloud_day),
                       ("filas_nuvosidad_tarde", config.rows_cloud_afternoon)):
        for row in rows:
            nub = _num(cell(sheet_cloud, row, config.col_nuvosidad))
            if nub is not None:
                horarias.append({"hora": hora, "datos": {"nub": nub}})

    rain_row, rain_col = config.rain_cell
    pp = _num(cell(config.sheet_rain, rain_row, rain_col))

    doc = {"meta": {"estacion": estacion, "fecha": fecha},
           "horarias": horarias, "resumen_dia": {"pp": pp}}

    year = fecha[4:8] if len(fecha) == 8 else "1970"
    month = fecha[2:4] if len(fecha) == 8 else "01"
    day = fecha[0:2] if len(fecha) == 8 else "01"
    text = json.dumps(doc, ensure_ascii=False)
    sha = _sha256_text(text)
    conn.execute("BEGIN IMMEDIATE")
    try:
        tx = _NoCommitConnection(conn)
        _ensure_station_for_import(tx, estacion)
        summaries_repo.upsert_summary_doc(
            tx, station_id=estacion, year=int(year), month=int(month),
            kind="diario", data=text, sha256=sha, day=int(day),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise

    obs = observation_from_doc(doc)
    if obs is None:  # pragma: no cover — doc built above always has meta
        raise SummaryError("Failed to read excel: empty meta", stage="import", path=str(path))
    return obs


def _ensure_station_for_import(conn: sqlite3.Connection, station: str) -> None:
    """Same placeholder contract as the F1 importer (unknown station →
    inactive row, verbatim id): foreign keys keep holding the diario doc."""
    if conn.execute("SELECT 1 FROM stations WHERE id = ?", (station,)).fetchone():
        return
    conn.execute(
        "INSERT INTO stations (id, name, provincia, latitud, longitud, elevacion, ch, is_active)"
        " VALUES (?, ?, '', 0.0, 0.0, 0.0, 0.0, 0)",
        (station, station),
    )


__all__ = [
    "HEADERS",
    "HOURS_8",
    "DailyObservation",
    "DaySummary",
    "ExcelConfig",
    "HourlyData",
    "HourlyRecord",
    "ObservationMeta",
    "SaveError",
    "SummaryError",
    "apply_corrections",
    "corrections_by_day",
    "delete_doc",
    "generate_month",
    "get_default_config",
    "get_diario_doc",
    "import_observation_excel",
    "list_history",
    "load_doc",
    "observation_from_doc",
]
