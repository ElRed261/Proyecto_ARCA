"""Excel export for the monthly summary (PRD F5).

Port of ``OLD/frontend/src-tauri/src/monthly_summary/json_pipeline.rs::

export_to_excel`` in openpyxl terms: ONE worksheet, the 21 verbatim oracle
headers (bold), numbers written as numbers, strings as strings, nulls as
empty cells. Qt-free (ADR-03). The dependency is pip-installed into the
running venv (parent-authorized F5 toolchain addition).
"""

from __future__ import annotations

import pathlib
from typing import Any

import openpyxl
from openpyxl.styles import Font

from oms_core.services.summary import HEADERS, INTEGER_HEADERS


def export_to_excel(doc: dict[str, Any], out_path: str | pathlib.Path) -> pathlib.Path:
    """Write the monthly summary doc to ``out_path`` and return it.

    Row 1 holds the header (bold, verbatim keys). Cells hold the exact value
    from the document: numbers stay numbers, direction columns stay strings,
    ``null`` cells are left empty.
    """
    out_path = pathlib.Path(out_path)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active

    bold = Font(bold=True)
    for column, header in enumerate(HEADERS, start=1):
        worksheet.cell(row=1, column=column, value=header).font = bold

    for row_index, kpis in enumerate(doc.get("datos") or [], start=2):
        for column, header in enumerate(HEADERS, start=1):
            value = kpis.get(header)
            if value is not None:
                cell = worksheet.cell(row=row_index, column=column, value=value)
                if header in INTEGER_HEADERS:
                    cell.number_format = "0"

    workbook.save(out_path)
    return out_path


__all__ = ["export_to_excel"]
