"""F5 — Excel export/import tests (openpyxl).

Covers: the 21 VERBATIM oracle headers + bold header row, numbers stored as
numbers / strings as strings / nulls as empty cells, the re-read round-trip,
and the xlsx import mirroring the oracle ``excel_reader.rs`` mapping (sheets
3074/4074/1200Z, config rows/cols, filename date/station extraction, missing
optional sheets tolerated) feeding a real generation.
"""

import json
from pathlib import Path

import pytest

pytest.importorskip("openpyxl", reason="openpyxl not installed")

import openpyxl

from oms_core.importer.importer import Importer
from oms_core.services import excel_export
from oms_core.services import summary as summary_service

HEADERS = list(summary_service.HEADERS)

# oracle ExcelConfig (excel_reader.rs port): 0-indexed rows/columns
ROWS_8 = [10, 13, 16, 19, 22, 25, 28, 31]
ROWS_CLOUD_DAY = [13, 16, 19]
ROWS_CLOUD_AFTERNOON = [22, 25, 28]
ROWS_4_TEMPERATURES = [10, 16, 22, 28]


def _sample_doc() -> dict:
    """A hand-built 2-day doc with numbers, strings and nulls."""
    return {
        "meta": {"estacion": "78485", "periodo": "10/2025", "total_dias": 2,
                 "generado": "2026-01-01T00:00:00-04:00"},
        "datos": [
            {
                "Dias": "2025-10-01", "Mayor temperatura máxima": 31.8,
                "Menor temperatura minima": 24.0, "Media temperatura": 27.9,
                "MAxima presión nivel medio del mar": 1015.7,
                "Minima presión nivel medio del mar": 1012.9,
                "Media presión nivel medio del mar": 1014.3,
                "Media presion en la estación": 1012.2, "Lluvia": None,
                "Dirección del viento": "SSE", "Velocidad media del viento": 7.1,
                "Velocidad máxima y direccion": "SSE 12,6",
                "Recorrido del viento": 170.4, "Nuvocidad dia": 3.0,
                "Nuvocidad noche": None, "Media de nuvocidad": 3.0,
                "Humedad maxima": 93.0, "Humedad minima": 80.0,
                "Humedad media": 86.5, "Punto de rocio": 23.1,
                "Tensión de vapor": 28.4,
            },
            {
                "Dias": "2025-10-02", "Mayor temperatura máxima": None,
                "Menor temperatura minima": None, "Media temperatura": None,
                "MAxima presión nivel medio del mar": None,
                "Minima presión nivel medio del mar": None,
                "Media presión nivel medio del mar": None,
                "Media presion en la estación": None, "Lluvia": 0.0,
                "Dirección del viento": None, "Velocidad media del viento": None,
                "Velocidad máxima y direccion": "N/A",
                "Recorrido del viento": None, "Nuvocidad dia": None,
                "Nuvocidad noche": None, "Media de nuvocidad": None,
                "Humedad maxima": None, "Humedad minima": None,
                "Humedad media": None, "Punto de rocio": None,
                "Tensión de vapor": None,
            },
        ],
    }


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------


def test_export_sanity_header_fingerprint(tmp_path):
    """The exact 21 headers byte sequence the oracle writes (order matters)."""
    assert HEADERS == [
        "Dias", "Mayor temperatura máxima", "Menor temperatura minima",
        "Media temperatura", "MAxima presión nivel medio del mar",
        "Minima presión nivel medio del mar", "Media presión nivel medio del mar",
        "Media presion en la estación", "Lluvia", "Dirección del viento",
        "Velocidad media del viento", "Velocidad máxima y direccion",
        "Recorrido del viento", "Nuvocidad dia", "Nuvocidad noche",
        "Media de nuvocidad", "Humedad maxima", "Humedad minima",
        "Humedad media", "Punto de rocio", "Tensión de vapor",
    ]


def test_export_writes_file_with_verbatim_bold_headers(tmp_path):
    out = tmp_path / "resumen.xlsx"

    result = excel_export.export_to_excel(_sample_doc(), out)

    assert result == out
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == HEADERS
    assert all(cell.font.bold for cell in ws[1])


def test_export_numbers_are_numbers_and_strings_are_strings(tmp_path):
    out = tmp_path / "resumen.xlsx"
    excel_export.export_to_excel(_sample_doc(), out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # Dias column: strings
    assert ws.cell(row=2, column=1).value == "2025-10-01"
    assert isinstance(ws.cell(row=2, column=1).value, str)
    # numeric kpis: numbers
    assert ws.cell(row=2, column=2).value == 31.8
    assert isinstance(ws.cell(row=2, column=2).value, (int, float))
    # direction columns: strings
    assert ws.cell(row=2, column=10).value == "SSE"
    assert isinstance(ws.cell(row=2, column=12).value, str)
    # null cells stay empty (None)
    assert ws.cell(row=2, column=9).value is None  # Lluvia null
    assert ws.cell(row=3, column=2).value is None


def test_export_from_generated_doc_roundtrip(fresh_db, tmp_path):
    Importer(fresh_db).import_root(
        Path(__file__).parent / "fixtures" / "importer" / "json")
    doc = summary_service.generate_month(fresh_db, "Estación Central", 2025, 5)
    out = tmp_path / "resumen_Estacion Central_052025.xlsx"

    result = excel_export.export_to_excel(doc, out)

    wb = openpyxl.load_workbook(result)
    ws = wb.active
    assert ws.max_row == 32  # header + 31 calendar days
    headers = [cell.value for cell in ws[1]]
    assert headers == HEADERS
    # day 15 carries the fixture day data
    kpis15 = doc["datos"][14]
    assert ws.cell(row=16, column=1).value == "2025-05-15"
    assert ws.cell(row=16, column=2).value == kpis15["Mayor temperatura máxima"]


# ---------------------------------------------------------------------------
# import (excel_reader.rs port)
# ---------------------------------------------------------------------------


def _build_excel_corpus(tmp_path: Path) -> Path:
    """An xlsx laid out exactly like the oracle config: 3074/4074/1200Z."""
    wb = openpyxl.Workbook()
    ws3074 = wb.active
    ws3074.title = "3074"
    ws3074.cell(row=3, column=2).value = "78485"  # B3 station

    for r in ROWS_8:
        # 0-indexed cols: presion_estacion 1, pres_nmm 2, punto_rocio 8,
        # t_vapor 9, hr 10, viento_dir 11, viento_vel 12
        ws3074.cell(row=r + 1, column=2).value = 1012.0
        ws3074.cell(row=r + 1, column=3).value = 1014.0
        ws3074.cell(row=r + 1, column=9).value = 22.0
        ws3074.cell(row=r + 1, column=10).value = 27.0
        ws3074.cell(row=r + 1, column=11).value = 80.0
        ws3074.cell(row=r + 1, column=12).value = "60"
        ws3074.cell(row=r + 1, column=13).value = 3.0

    ws4074 = wb.create_sheet("4074")
    for r, nub in zip(ROWS_CLOUD_DAY, [7.0, 6.0, 5.0]):
        ws4074.cell(row=r + 1, column=2).value = nub  # col_nuvosidad 1
    for r, nub in zip(ROWS_CLOUD_AFTERNOON, [2.0, 6.0, 5.0]):
        ws4074.cell(row=r + 1, column=2).value = nub
    for r in ROWS_4_TEMPERATURES:
        ws4074.cell(row=r + 1, column=22).value = 30.0  # col_temp_maxima 21
        ws4074.cell(row=r + 1, column=23).value = 22.0  # col_temp_minima 22

    ws1200 = wb.create_sheet("1200Z")
    ws1200.cell(row=37, column=7).value = 4.2  # G37 rain cell (0-idx 36,6)

    path = tmp_path / "78485 08102025.xlsx"
    wb.save(path)
    return path


def test_import_observation_excel_maps_sheets_and_config(fresh_db, tmp_path):
    path = _build_excel_corpus(tmp_path)
    obs = summary_service.import_observation_excel(fresh_db, path)

    assert isinstance(obs, summary_service.DailyObservation)
    assert obs.meta.estacion == "78485"  # B3 wins over the filename
    assert obs.meta.fecha == "08102025"
    horas = {row.hora: row.datos for row in obs.horarias}
    assert set(horas) == {
        "06Z", "09Z", "12Z", "15Z", "18Z", "21Z", "00Z", "03Z",
        "TMAX", "TMIN", "filas_nuvosidad_dia", "filas_nuvosidad_tarde",
    }
    h06 = horas["06Z"]
    assert h06.p_est == 1012.0 and h06.p_nmm == 1014.0
    assert h06.rocio == 22.0 and h06.t_vapor == 27.0
    assert h06.hr == 80.0
    assert h06.v_dir_deg == 60.0  # "60" parsed from a string cell
    assert h06.v_vel == 3.0
    tmax_vals = [row.datos.tmax for row in obs.horarias if row.hora == "TMAX"]
    assert tmax_vals == [30.0, 30.0, 30.0, 30.0]
    nub_tarde = [
        row.datos.nub for row in obs.horarias if row.hora == "filas_nuvosidad_tarde"
    ]
    assert nub_tarde == [2.0, 6.0, 5.0]
    assert obs.resumen_dia.pp == 4.2

    # the imported day document is stored as a diario doc (verbatim shape)
    stored = json.loads(
        summary_service.get_diario_doc(fresh_db, "78485", 2025, 10, 8))
    assert stored["meta"]["fecha"] == "08102025"
    assert stored["resumen_dia"]["pp"] == 4.2


def test_imported_excel_day_flows_into_generation(fresh_db, tmp_path):
    path = _build_excel_corpus(tmp_path)
    summary_service.import_observation_excel(fresh_db, path)

    doc = summary_service.generate_month(fresh_db, "78485", 2025, 10)
    kpis = doc["datos"][7]
    assert kpis["Dias"] == "2025-10-08"
    assert kpis["Lluvia"] == 4.2  # own-day rain (product decision)
    assert kpis["Nuvocidad dia"] == 6.0  # (7+6+5)/3
    assert kpis["Nuvocidad noche"] == 4.3  # (2+6+5)/3 → 4.333 → half-up 4.3
    # media from unrounded dia/tarde halves: (6.0 + 4.333…) / 2 → 5.2
    assert kpis["Media de nuvocidad"] == 5.2
    assert kpis["Mayor temperatura máxima"] == 30.0
    assert kpis["Velocidad media del viento"] == 10.8  # 3.0 m/s → km/h


def test_import_excel_missing_optional_sheets_tolerated(fresh_db, tmp_path):
    """4074/1200Z missing → empty ranges (oracle unwrap_or_else(empty))."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3074"
    ws.cell(row=3, column=2).value = ""  # empty B3 → filename fallback
    path = tmp_path / "Estación X 01012025.xlsx"
    wb.save(path)

    obs = summary_service.import_observation_excel(fresh_db, path)
    assert obs.meta.estacion == "Estación X"  # from the filename
    assert obs.meta.fecha == "01012025"
    rows = [row for row in obs.horarias if row.hora in summary_service.HOURS_8]
    assert len(rows) == 8
    assert all(v is None for row in rows for v in vars(row.datos).values())
    assert obs.resumen_dia.pp is None


def test_humidity_extremes_export_as_whole_numbers(tmp_path):
    """Owner's rule (ronda 32): HR max/min export as integers, mean with one
    decimal. The cell keeps the number; the format is what makes it whole."""
    from openpyxl import load_workbook

    from oms_core.services.excel_export import export_to_excel
    from oms_core.services.summary import HEADERS

    doc = {"meta": {"estacion": "78485", "periodo": "10/2025"},
           "datos": [{"Dias": "2025-10-01", "Humedad maxima": 97.0,
                      "Humedad minima": 66.0, "Humedad media": 84.8}]}
    out = export_to_excel(doc, tmp_path / "hum.xlsx")
    sheet = load_workbook(out).active
    hmax = HEADERS.index("Humedad maxima") + 1
    hmin = HEADERS.index("Humedad minima") + 1
    hmed = HEADERS.index("Humedad media") + 1
    assert sheet.cell(row=2, column=hmax).number_format == "0"
    assert sheet.cell(row=2, column=hmin).number_format == "0"
    assert sheet.cell(row=2, column=hmed).number_format == "General"
    assert sheet.cell(row=2, column=hmax).value == 97.0
