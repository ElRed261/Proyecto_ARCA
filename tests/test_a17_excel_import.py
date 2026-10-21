"""A17 — Excel (.xlsm/.xlsx) import into the schema 2 document.

Ported from the owner's original transformer; these tests build the workbooks
with openpyxl so the cell layout stays visible in the test itself.
"""

import pytest
from openpyxl import Workbook

from oms_core.services import excel_import as import_mod

STATION = "78486"


def _z_hour_sheet(wb, name, *, yyggiw=None, observer="YASSER GALAN",
                  data_cells=None, groups=None):
    ws = wb.create_sheet(name)
    if observer:
        ws.cell(row=6, column=2, value=observer)
    ws.cell(row=8, column=2, value=yyggiw or name[:4])
    for (row, col), value in (data_cells or {}).items():
        ws.cell(row=row, column=col, value=value)
    for (row, col), value in (groups or {}).items():
        ws.cell(row=row, column=col, value=value)
    return ws


def _new_workbook(tmp_path, nombre="06062024.xlsm"):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(
        wb, "0006Z", yyggiw="06061",
        data_cells={(29, 1): 28.4, (31, 3): 1012.4, (29, 6): 31.2},
        groups={(10, 1): "32566", (10, 2): "50602", (10, 7): "70311"},
    )
    sheet = wb.create_sheet("3074")
    sheet.cell(row=11, column=2, value=1)          # hour 1 -> form row 2
    sheet.cell(row=11, column=3, value=1012.0)
    sheet.cell(row=11, column=5, value=1012.1)     # pres_alti (kept: manual)
    sheet.cell(row=11, column=14, value=4)         # viento_vel (kept: manual)
    sheet.cell(row=11, column=16, value="03")      # tiempo_presente
    sheet.cell(row=11, column=17, value=1)         # granizo
    path = tmp_path / nombre
    wb.save(path)
    wb.close()
    return path


# ------------------------------------------------------------ format + meta


def test_the_new_layout_is_detected(tmp_path):
    from openpyxl import load_workbook

    wb = load_workbook(_new_workbook(tmp_path))
    try:
        assert import_mod.detect_format(wb) == "new"
    finally:
        wb.close()


def test_the_old_layout_is_detected_when_there_is_no_3074_sheet(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061")
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    from openpyxl import load_workbook

    wb2 = load_workbook(path)
    try:
        assert import_mod.detect_format(wb2) == "old"
    finally:
        wb2.close()


def test_the_split_digits_layout_is_detected(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061")
    sheet = wb.create_sheet("3074")
    sheet.cell(row=11, column=3, value="0")
    sheet.cell(row=11, column=4, value="1")
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    from openpyxl import load_workbook

    wb2 = load_workbook(path)
    try:
        assert import_mod.detect_format(wb2) == "horrible"
    finally:
        wb2.close()


def test_the_date_comes_from_the_file_name_and_reads_as_iso(tmp_path):
    document = import_mod.parse_excel(_new_workbook(tmp_path), station=STATION)
    assert document["schema"] == 2
    assert document["meta"]["fecha"] == "2024-06-06"
    assert document["meta"]["estacion"] == STATION
    assert document["meta"]["actualizado"]


def test_the_station_is_read_from_the_sheet_when_not_given(tmp_path):
    path = _new_workbook(tmp_path)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb["0006Z"].cell(row=8, column=3, value=STATION)
    wb.save(path)
    wb.close()
    assert import_mod.parse_excel(path)["meta"]["estacion"] == STATION


def test_the_date_override_wins_over_the_file_name(tmp_path):
    document = import_mod.parse_excel(
        _new_workbook(tmp_path), station=STATION, date="2024-12-31"
    )
    assert document["meta"]["fecha"] == "2024-12-31"


# --------------------------------------------------------------- the hours


def test_each_hour_keeps_its_observer(tmp_path):
    document = import_mod.parse_excel(_new_workbook(tmp_path), station=STATION)
    assert document["synop"]["06Z"]["observador"] == "YASSER GALAN"


def test_the_observed_values_land_on_their_form_keys(tmp_path):
    hour = import_mod.parse_excel(
        _new_workbook(tmp_path), station=STATION
    )["synop"]["06Z"]
    assert hour["datos"]["ts"] == "28.4"
    assert hour["datos"]["Tmax"] == "31.2"
    assert hour["datos"]["pres_est"] == "1012.4"


def test_a_short_pressure_gains_its_thousand(tmp_path):
    """The station types 012.4; the thousand is implied by the layout."""
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061", data_cells={(31, 3): 12.4})
    _z_hour_sheet(wb, "0012Z", yyggiw="12061")
    wb.create_sheet("3074")  # the 'new' layout needs the sheet to exist
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    hour = import_mod.parse_excel(path, station=STATION)["synop"]["06Z"]
    assert hour["datos"]["pres_est"] == "1012.4"


def test_the_groups_keep_their_wmo_names(tmp_path):
    hour = import_mod.parse_excel(
        _new_workbook(tmp_path), station=STATION
    )["synop"]["06Z"]
    assert hour["grupos"]["Nddff"] == "50602"
    assert hour["grupos"]["7wwW1W2"] == "70311"


def test_a_placeholder_group_is_not_stored(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061", groups={(10, 2): "9"})
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    hour = import_mod.parse_excel(path, station=STATION)["synop"]["06Z"]
    assert "Nddff" not in hour["grupos"]


# ------------------------------------------------------------- the CLI 3074


def test_the_3074_rows_are_keyed_by_their_form_row(tmp_path):
    cli = import_mod.parse_excel(
        _new_workbook(tmp_path), station=STATION
    )["cli"]["3074"]
    assert list(cli["horas"]) == ["2"], "hour 1 of the sheet is form row 2"


def test_the_3074_keeps_only_what_the_operator_types(tmp_path):
    row = import_mod.parse_excel(
        _new_workbook(tmp_path), station=STATION
    )["cli"]["3074"]["horas"]["2"]
    assert row["viento_vel"] == "4"
    assert row["pres_alti"] == "1012.1"
    assert row["fenomenos"]["granizo"] is True
    assert row["fenomenos"]["tiempo_presente"] == "03"
    # derived values are not stored (rule R2)
    assert "pres_est" not in row
    assert "visibilidad" not in row
    # unchecked phenomena never survive
    assert "trueno" not in row["fenomenos"]


def test_the_split_digits_layout_rebuilds_the_pressure(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061")
    sheet = wb.create_sheet("3074")
    sheet.cell(row=11, column=2, value=1)
    for column, digit in zip((3, 4, 5, 6), "0120", strict=True):
        sheet.cell(row=11, column=column, value=digit)
    # the altimeter is one of the few values the operator types, so it is the
    # one that survives the cleaning of the CLI block
    for column, digit in zip((11, 12, 13, 14), "0121", strict=True):
        sheet.cell(row=11, column=column, value=digit)
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    row = import_mod.parse_excel(path, station=STATION)["cli"]["3074"]["horas"]["2"]
    assert row["pres_alti"] == "1012.1"  # decoded from the split digits


# ------------------------------------------------------------------- errors


def test_a_book_without_hours_is_refused(tmp_path):
    wb = Workbook()
    wb.active.title = "Hoja1"
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="no reconocido"):
        import_mod.parse_excel(path, station=STATION)


def test_a_book_without_a_date_is_refused(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061")
    path = tmp_path / "sin_fecha.xlsx"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="fecha"):
        import_mod.parse_excel(path, station=STATION)


def test_a_book_without_a_station_is_refused(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    _z_hour_sheet(wb, "0006Z", yyggiw="06061")
    path = tmp_path / "06062024.xlsx"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="estación"):
        import_mod.parse_excel(path)


# ------------------------------------------------- the picker applies a file


def _picker(conn, qtbot, station="78486"):
    from oms_ui.observation_picker import ObservationPickerDialog

    dialog = ObservationPickerDialog(conn, station=station)
    qtbot.addWidget(dialog)
    return dialog


def test_the_picker_stores_the_day_of_the_workbook(fresh_db, qtbot, tmp_path):
    """Importing writes the day: hours through the app's services, and the
    observer with them."""
    from oms_core.services.document import load_day_document

    path = _new_workbook(tmp_path)
    dialog = _picker(fresh_db, qtbot)
    assert dialog._import_from_path(str(path)) is True
    assert dialog.selection() == (STATION, "2024-06-06")
    day = load_day_document(fresh_db, STATION, "2024-06-06")
    assert day.hour("06Z")["ts"] == "28.4"
    assert day.observer("06Z") == "YASSER GALAN"
    assert day.cli["cli3074"]["horas"]["2"]["pres_alti"] == "1012.1"


def test_an_existing_day_asks_before_replacing(fresh_db, qtbot, tmp_path, monkeypatch):
    """Owner's rule: never overwrite silently."""
    from oms_core.services import synoptic as synoptic_service
    from oms_ui import observation_picker as ui

    synoptic_service.save_hour(fresh_db, STATION, "2024-06-06", "06Z",
                               {"ts": "11.1"})
    asked: dict[str, object] = {}

    def fake_question(*args, **kwargs):
        asked["asked"] = True
        return ui.QMessageBox.StandardButton.Cancel

    monkeypatch.setattr(ui.QMessageBox, "question", staticmethod(fake_question))
    dialog = _picker(fresh_db, qtbot)
    assert dialog._import_from_path(str(_new_workbook(tmp_path))) is False
    assert asked.get("asked") is True
    # the stored day is untouched
    from oms_core.services.document import load_day_document

    assert load_day_document(fresh_db, STATION, "2024-06-06").hour("06Z")["ts"] == "11.1"


def test_an_existing_day_is_replaced_when_confirmed(fresh_db, qtbot, tmp_path, monkeypatch):
    from oms_core.services import synoptic as synoptic_service
    from oms_core.services.document import load_day_document
    from oms_ui import observation_picker as ui

    synoptic_service.save_hour(fresh_db, STATION, "2024-06-06", "06Z",
                               {"ts": "11.1"})
    monkeypatch.setattr(
        ui.QMessageBox, "question",
        staticmethod(lambda *a, **k: ui.QMessageBox.StandardButton.Yes),
    )
    dialog = _picker(fresh_db, qtbot)
    assert dialog._import_from_path(str(_new_workbook(tmp_path))) is True
    assert load_day_document(fresh_db, STATION, "2024-06-06").hour("06Z")["ts"] == "28.4"


def test_a_broken_workbook_is_reported_and_changes_nothing(fresh_db, qtbot, tmp_path, monkeypatch):
    from oms_ui import observation_picker as ui

    warned: dict[str, object] = {}

    def fake_warning(*args, **kwargs):
        warned["message"] = args[2]

    monkeypatch.setattr(ui.QMessageBox, "warning", staticmethod(fake_warning))
    bad = tmp_path / "sin_fecha.xlsx"
    wb = Workbook()
    wb.active.title = "Hoja1"
    wb.save(bad)
    dialog = _picker(fresh_db, qtbot)
    assert dialog._import_from_path(str(bad)) is False
    assert "no reconocido" in str(warned.get("message"))
    assert dialog.selected_fecha is None
