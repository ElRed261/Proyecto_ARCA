"""F5 — pytest-qt UI tests for the summary page.

QT_QPA_PLATFORM=offscreen is forced in tests/conftest.py; the real flow is
exercised with clicks: station/month/year selectors → Generar → the 31×21
table (QTableView over a QAbstractTableModel) with verbatim Spanish headers
and the corrections indicator → export to Excel writes a readable file.
"""

from pathlib import Path

import pytest

pytest.importorskip("PySide6", reason="PySide6 not installed in this interpreter")

from PySide6.QtCore import Qt

from oms_core.db.repositories import audit as audit_repo
from oms_core.importer.importer import Importer
from oms_core.services import summary as summary_service
from oms_ui.pages.summary import SummaryPage

FIXTURE_JSON = Path(__file__).parent / "fixtures" / "importer" / "json"


@pytest.fixture()
def page(qtbot, fresh_db):
    Importer(fresh_db).import_root(FIXTURE_JSON)
    widget = SummaryPage(fresh_db)
    qtbot.addWidget(widget)
    return widget


# ---------------------------------------------------------------- selectors


def test_selectors_list_stations_months_years(page):
    ids = [page.station_combo.itemText(i) for i in range(page.station_combo.count())]
    assert "Estación Central" in ids  # from the imported fixture diarios
    assert page.month_combo.count() == 12
    assert page.month_combo.itemText(0) == "enero"
    assert page.month_combo.itemText(4) == "mayo"
    assert int(page.year_combo.currentText()) >= 2025


# ---------------------------------------------------------------- full flow


def test_generate_flow_31_rows_and_headers(page, qtbot, tmp_path):
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)  # mayo
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)

    model = page.table.model()
    assert model.rowCount() == 31
    assert model.columnCount() == 21
    assert model.headerData(0, Qt.Horizontal) == "Dias"
    assert model.headerData(20, Qt.Horizontal) == "Tensión de vapor"
    # the day with data shows the oracle value; a day without the diario is empty
    assert model.data(model.index(14, 0), Qt.DisplayRole) == "2025-05-15"
    assert model.data(model.index(14, 1), Qt.DisplayRole) == "29.0"
    assert model.data(model.index(9, 1), Qt.DisplayRole) == ""
    # the generated doc is stored and the history reflects it
    assert summary_service.load_doc(
        page.conn, "Estación Central", 2025, 5) is not None
    assert len(page.history_entries) == 1
    assert page.history_entries[0]["period"] == "05/2025"
    assert len(page.history_cards) == 1


def test_history_click_loads_doc_into_table(page, qtbot):
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)

    page.open_history_entry(0)  # what clicking the card does
    model = page.table.model()
    assert model.rowCount() == 31


def test_export_button_writes_excel(page, qtbot, tmp_path):
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)

    out = tmp_path / "exportada.xlsx"
    page.export_doc_to(out)  # the handler behind the export button
    import openpyxl

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert [c.value for c in ws[1]] == list(summary_service.HEADERS)
    assert ws.cell(row=16, column=1).value == "2025-05-15"


def test_corrections_indicator_marks_the_day(page, qtbot):
    audit_repo.propose_correction(
        page.conn,
        station_id="Estación Central",
        fecha="2025-05-15",
        hora="12Z",
        campo="pres_est",
        valor_original="1013.0",
        valor_corregido="1020.0",
        justificacion="test",
        corregido_por="tester",
        estado="aprobado",
    )
    page.include_corrections.setChecked(True)  # the overlay marks days
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)

    editor = page.table.model()
    idx = editor.createIndex(14, 0)
    assert editor.data(idx, Qt.BackgroundRole) is not None  # marked day
    assert "1" in str(editor.data(idx, Qt.ToolTipRole))


def test_generate_month_without_data_shows_null_table(page, qtbot):
    page.station_combo.setCurrentText("MDJB")
    page.month_combo.setCurrentIndex(1)  # febrero
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)

    model = page.table.model()
    assert model.rowCount() == 28
    assert model.data(model.index(0, 1), Qt.DisplayRole) == ""
    # the doc sort of exists: 28 rows of nulls
    assert len(summary_service.load_doc(page.conn, "MDJB", 2025, 2)["datos"]) == 28


# ------------------------------------------------- ronda 13: rediseño 1:1


def _texts(widget):
    from PySide6.QtWidgets import QLabel

    return [label.text() for label in widget.findChildren(QLabel)]


def test_summary_speaks_the_image_verbatim(page):
    """Title, subtitle, panel heading and the empty state copy are the
    owner's strings, character for character (image = brief)."""
    texts = _texts(page)
    assert "Resumen Mensual SYNOP" in texts
    assert "Generación y visualización de resúmenes climáticos" in texts
    assert "CARGA AUTOMÁTICA POR ESTACIÓN Y PERÍODO" in texts
    assert "Historial de Reportes Guardados" in texts
    assert "No hay resúmenes" in texts
    assert (
        "Selecciona una estación y un período arriba para generar tu "
        "primer resumen mensual, o importa archivos manualmente."
    ) in texts
    assert page.generate_button.text() == "Cargar Mes"
    assert page.manual_button.text() == "Generar Manual"
    assert page.history_tab_button.text() == "Historial"
    assert page.include_corrections.text() == "Incluir correcciones (overlay)"


def test_tabla_pill_needs_a_document(page, qtbot):
    """The image shows only 'Historial'; the table pill appears with a doc."""
    # isHidden() is the local flag: isVisible() lies when the page
    # itself was never shown (offscreen tests).
    assert page.table_tab_button.isHidden()
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)
    assert not page.table_tab_button.isHidden()
    assert page.stack.currentIndex() == 1  # the oracle opens the table
    qtbot.mouseClick(page.history_tab_button, Qt.LeftButton)
    assert page.stack.currentIndex() == 0


def test_empty_state_yields_to_cards(page, qtbot):
    assert not page.empty_state.isHidden()
    assert not page.history_cards
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)
    assert page.empty_state.isHidden()
    assert len(page.history_cards) == 1


def test_overlay_switch_changes_what_the_table_shows(page, qtbot):
    """S8: the checkbox is real — off generates the month untouched and
    without correction marks; on applies the approved overlay."""
    audit_repo.propose_correction(
        page.conn,
        station_id="Estación Central",
        fecha="2025-05-15",
        hora="12Z",
        campo="pres_est",
        valor_original="1013.0",
        valor_corregido="1020.0",
        justificacion="test",
        corregido_por="tester",
        estado="aprobado",
    )
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)
    page.year_combo.setCurrentText("2025")

    def row_values():
        model = page.table.model()
        return [model.data(model.index(14, c), Qt.DisplayRole) for c in range(21)]

    def day_mark():
        model = page.table.model()
        return model.data(model.index(14, 0), Qt.BackgroundRole)

    page.include_corrections.setChecked(False)
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)
    without = row_values()
    assert day_mark() is None

    page.include_corrections.setChecked(True)
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)
    with_overlay = row_values()
    assert day_mark() is not None
    assert without != with_overlay  # the correction really reached the doc


def test_page_flows_in_a_single_scroll_region(page):
    """The legacy double-scroll bug must not come back: one page scroll."""
    from PySide6.QtWidgets import QScrollArea

    assert len(page.findChildren(QScrollArea)) == 1


# ------------------------------------------- humidity presentation (ronda 32)


def test_humidity_extremes_are_integers_and_the_mean_keeps_one_decimal(page, qtbot):
    """Owner's rule: HR max/min are whole numbers, the monthly mean has one
    decimal. The calculation already rounds the extremes to 0 decimals; this
    pins the presentation so a stored 97.0 never reads as "97.0"."""
    page.station_combo.setCurrentText("Estación Central")
    page.month_combo.setCurrentIndex(4)  # mayo
    page.year_combo.setCurrentText("2025")
    qtbot.mouseClick(page.generate_button, Qt.LeftButton)
    model = page.table.model()

    hmax, hmin, hmed = 16, 17, 18  # HEADERS indices
    for column in (hmax, hmin):
        for row in range(model.rowCount()):
            text = model.data(model.index(row, column), Qt.DisplayRole)
            if text:
                assert "." not in text, f"columna {column}: {text!r} no es entero"
    seen = False
    for row in range(model.rowCount()):
        text = model.data(model.index(row, hmed), Qt.DisplayRole)
        if text:
            seen = True
            assert len(text.split(".")[1]) == 1, f"media {text!r} sin un decimal"
    assert seen, "el fixture debe traer humedad para probar el formato"


def test_the_mean_of_rounded_humidity_keeps_its_decimal():
    """96.8 and 97.0 both display as integers, but the mean is untouched."""
    assert summary_service.display_value("Humedad maxima", 96.8) == "97"
    assert summary_service.display_value("Humedad minima", 97.0) == "97"
    assert summary_service.display_value("Humedad media", 84.75) == "84.75"
    assert summary_service.display_value("Humedad maxima", None) == ""
    assert summary_service.display_value("Media temperatura", 27.0) == "27.0"
